"""Step 4 fix pass: client status vs the audited PDF, removed public report route, recovery lookup, key rotation,
edit-vs-submit race, render-time cleaning, wipe, client visibility."""
import io
import os
import sys

import pytest
from docx import Document

sys.path.insert(0, os.path.dirname(__file__))
from audit_helpers import PW, world  # noqa: E402,F401
import db  # noqa: E402
import reportcontent  # noqa: E402
import reportdoc  # noqa: E402
import reportrender  # noqa: E402


def _docx_text(data: bytes) -> str:
    d = Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        parts += [c.text for r in t.rows for c in r.cells]
    return "\n".join(parts)


def _stage(tid, stage):
    db.get_conn().execute("UPDATE proposals SET stage=? WHERE id=?", (stage, tid))
    db.get_conn().commit()


def _render(w):
    t = db.get_proposal(w.tid, org_id=None)
    c = reportcontent.ensure_v1(t)
    return reportrender.render_docx(c["content"], t, "PT A", reportdoc.findings_for(t))


# --- 1. client status ---
def test_report_does_not_print_the_client_status(monkeypatch, tmp_path):
    w = world(monkeypatch, tmp_path)
    before = _docx_text(_render(w))
    assert "Status" not in before
    t = db.get_proposal(w.tid, org_id=None)
    dg = reportdoc.digest(reportdoc.findings_for(t))
    db.set_finding(w.fids["SQL injection"], status="fixed")
    assert "Fixed" not in _docx_text(_render(w))
    assert reportdoc.digest(reportdoc.findings_for(t)) == dg


def test_client_status_is_blocked_until_delivery(api, monkeypatch, tmp_path):
    w = world(monkeypatch, tmp_path)
    tok = api.login("alice", PW)
    url = f"/api/findings/{w.fids['SQL injection']}/status"
    for stage in ("completed", "review_lead_pentester", "review_manager"):
        _stage(w.tid, stage)
        r = api.post(url, tok, {"status": "fixed"})
        assert r.status_code == 409 and "after the report is delivered" in r.text, stage
        assert db.get_finding(w.fids["SQL injection"], org_id=None)["status"] == "open"
    _stage(w.tid, "delivered")
    assert api.post(url, tok, {"status": "fixed"}).status_code == 200
    assert db.get_finding(w.fids["SQL injection"], org_id=None)["status"] == "fixed"


# --- 2. no unreviewed report on the public plane ---
def test_public_plane_has_no_legacy_report_routes(api, monkeypatch, tmp_path):
    w = world(monkeypatch, tmp_path)
    tok = api.login("alice", PW)
    paths = api.app.openapi()["paths"]
    assert "/api/scans/{job_id}/report" not in paths and "/api/reports/{fname}/download" not in paths
    assert api.post("/api/scans/j1/report", tok).status_code in (404, 405)
    assert api.get("/api/reports/anything.docx/download", tok).status_code in (404, 405)
