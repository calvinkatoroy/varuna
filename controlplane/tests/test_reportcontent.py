"""Content v1, the resolved view of content + live findings, and the docx renderer."""
import io
import os
import re
import sys

import pytest
from docx import Document

HERE = os.path.dirname(__file__)
sys.path.insert(0, HERE)
sys.path.insert(0, os.path.join(HERE, "..", "common"))
sys.path.insert(0, os.path.join(HERE, "..", "report"))
import db  # noqa: E402
import reportcontent  # noqa: E402
import reportdoc  # noqa: E402
import reportrender  # noqa: E402
from conftest import make_client, window  # noqa: E402

FINDINGS = [
    {"name": "SQL injection", "severity": "critical", "host": "8.8.8.8", "url": "/a?id=1", "evidence": "e1",
     "impact": "Data theft.", "remediation": "Use bound parameters."},
    {"name": "Missing header", "severity": "low", "host": "8.8.8.8", "url": "/", "evidence": "e2",
     "impact": "Minor.", "remediation": "Add the header."},
    {"name": "Noise", "severity": "medium", "host": "8.8.8.8", "url": "/n", "evidence": "e3",
     "impact": "None.", "remediation": "Ignore."},
]


def _task():
    org = make_client("alice", "PT A")
    nb, na = window()
    tid = db.create_proposal({"submitter": "alice", "org_id": org, "target": "http://8.8.8.8", "stage": "completed",
                              "assignee": "rizky", "job_id": "j1", "not_before": nb, "not_after": na,
                              "notes": "contact the NOC", "path": "/app", "port": 8080})
    db.save_findings("j1", "alice", org, FINDINGS)
    return db.get_proposal(tid, org_id=None)


def _text(docx_bytes):
    d = Document(io.BytesIO(docx_bytes))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts += [c.text for c in row.cells]
    return "\n".join(parts)


def test_baseline_has_the_seven_fixed_sections_and_locks_the_disclaimer():
    t = _task()
    c = reportcontent.baseline(t, "PT A")
    assert [s["id"] for s in c["sections"]] == ["cover", "executive_summary", "scope", "methodology",
                                                 "findings_summary", "findings", "conclusion"]
    assert [s["id"] for s in c["sections"] if s.get("ai_insert")] == ["executive_summary", "methodology", "conclusion"]
    scope = next(s for s in c["sections"] if s["id"] == "scope")
    locked = [b for b in scope["blocks"] if b["type"] == "paragraph"]
    assert len(locked) == 3 and not any(b["editable"] for b in locked)
    ids = [b["id"] for s in c["sections"] for b in s["blocks"]]
    assert len(ids) == len(set(ids))
    assert any(b["type"] == "finding_overrides" and b["overrides"] == {} for s in c["sections"] for b in s["blocks"])
    for s in c["sections"]:
        for b in s["blocks"]:
            if b["type"] in ("paragraph", "bullet") and b["editable"]:
                assert not re.search(r"\b\d+ (finding|critical|high)", b["text"])   # counts are computed, never typed


def test_ensure_v1_builds_once():
    t = _task()
    a = reportcontent.ensure_v1(t)
    b = reportcontent.ensure_v1(t)
    assert a["version"] == b["version"] == 1 and a["created_by"] == "system"
    assert len(db.list_content_versions(t["id"])) == 1


def test_resolve_follows_the_live_findings():
    t = _task()
    c = reportcontent.ensure_v1(t)["content"]
    findings = lambda: reportdoc.findings_for(t)

    def items(kind):
        return [i for s in reportrender.resolve(c, t, "PT A", findings()) for i in s["items"] if i["kind"] == kind]
    fs = items("finding")
    assert [(f["f"]["_id"], f["f"]["name"]) for f in fs] == [("VAR-001", "SQL injection"), ("VAR-002", "Noise"),
                                                              ("VAR-003", "Missing header")]
    assert items("counts")[0]["counts"]["critical"] == 1 and items("counts")[0]["risk"] == "Critical"
    noise = next(f for f in db.get_findings("j1") if f["name"] == "Noise")
    db.set_finding(noise["id"], verdict="fp")                                  # no content edit needed
    assert [f["f"]["name"] for f in items("finding")] == ["SQL injection", "Missing header"]
    assert items("counts")[0]["counts"]["medium"] == 0
    sql = next(f for f in db.get_findings("j1") if f["name"] == "SQL injection")
    ov = next(b for s in c["sections"] for b in s["blocks"] if b["type"] == "finding_overrides")
    ov["overrides"][sql["id"]] = {"remediation": "Rewritten advice."}
    assert items("finding")[0]["f"]["remediation"] == "Rewritten advice."
    assert items("finding")[1]["f"]["remediation"] == "Add the header."


def test_docx_contains_the_content_and_only_confirmed_findings():
    t = _task()
    c = reportcontent.ensure_v1(t)["content"]
    noise = next(f for f in db.get_findings("j1") if f["name"] == "Noise")
    db.set_finding(noise["id"], verdict="fp")
    out = reportrender.render_docx(c, t, "PT A", reportdoc.findings_for(t))
    assert out[:2] == b"PK"
    text = _text(out)
    assert "SQL injection" in text and "Use bound parameters." in text and "Missing header" in text
    assert "Noise" not in text and "VAR-001" in text
    assert "Executive Summary" in text and "Conclusion" in text
    assert "NOT assessed by this automated scan" in text      # the coverage disclaimer is always present


def test_hostile_text_renders_as_literal_text_and_never_breaks_the_document():
    t = _task()
    c = reportcontent.ensure_v1(t)["content"]
    bad = db.get_findings("j1")[0]
    db.set_finding(bad["id"], remediation="<script>alert(1)</script> \x1b[31mred\x00 & <b>x</b>")
    out = reportrender.render_docx(c, t, "PT A <i>", reportdoc.findings_for(t))
    text = _text(out)
    assert "<script>alert(1)</script>" in text and "\x1b" not in text and "\x00" not in text


def test_preview_is_json_safe_and_trims_evidence():
    import json
    t = _task()
    c = reportcontent.ensure_v1(t)["content"]
    db.set_finding(db.get_findings("j1")[0]["id"], evidence="x" * 5000)
    p = reportrender.preview(c, t, "PT A", reportdoc.findings_for(t))
    json.dumps(p)
    fs = [i["f"] for s in p for i in s["items"] if i["kind"] == "finding"]
    assert all(len(f["evidence"]) <= 500 for f in fs)
    assert set(fs[0]) == {"id", "_id", "name", "severity", "host", "url", "impact", "remediation", "description",
                          "cve", "cwe", "tool", "verdict", "evidence"}
