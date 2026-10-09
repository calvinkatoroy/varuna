"""Reports must survive hostile/dirty text: terminal escapes, control characters, RTL marks,
markup, absurd lengths. python-docx raises on XML-illegal characters, which used to stop every
template (and with it the automatic review) on a single ANSI code in scanner output."""
import io
import os
import sys

HERE = os.path.dirname(__file__)
sys.path.insert(0, os.path.join(HERE, "..", "common"))
sys.path.insert(0, os.path.join(HERE, "..", "report"))

from docx import Document  # noqa: E402

import generator  # noqa: E402

JOB = {"id": "j-hostile", "target": "http://t.example", "submitter": "acme\x00corp"}
DIRTY = [
    {"name": "SQLi \x1b[31mred\x1b[0m", "severity": "critical", "host": "h\x0bost", "url": "/a?x=\x00\x01",
     "evidence": "line1\x00\x08line2\x1b[0m emoji \U0001F600 rtl ‮ end", "description": "d\x0c",
     "impact": "<script>alert(1)</script> & \"quotes\" ]]>", "remediation": "fix\x7f", "cwe": "CWE-89",
     "cvss": 9.8, "tool": "sqlmap", "verdict": "tp"},
    {"name": "", "severity": "not-a-severity", "host": None, "url": None, "evidence": None, "tool": None},
    {"name": "N" * 20000, "severity": "low", "host": "h", "evidence": "E" * 200000},
]


def _text(data: bytes) -> str:
    d = Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts += [c.text for c in row.cells]
    return "\n".join(parts)


def test_every_template_survives_dirty_content():
    for template in generator.TEMPLATES:
        data = generator.generate(JOB, DIRTY, template)
        assert data[:2] == b"PK", f"{template} did not produce a docx"


def test_escapes_and_control_characters_are_removed_not_left_as_junk():
    text = _text(generator.generate(JOB, DIRTY, "Full Technical"))
    assert "\x1b" not in text and "\x00" not in text and "[31m" not in text
    assert "SQLi red" in text                         # the words survive, the colour codes do not
    assert "<script>alert(1)</script>" in text        # markup is shown as text, never interpreted


def test_absurdly_long_fields_are_truncated_not_dumped_into_the_report():
    text = _text(generator.generate(JOB, DIRTY, "Full Technical"))
    assert "[truncated]" in text
    assert len(text) < 60000, "a 200 KB evidence blob must not balloon the document"


def _task_for(tmp_path, job_id, **extra):
    import db
    db.reset_for_test(str(tmp_path / "t.db"))
    pid = db.create_proposal({"submitter": "acme", "target": "http://t.example", "org_id": "org-acme",
                              "not_before": "2026-10-09T08:00:00+00:00", "not_after": "2026-10-09T16:30:00+00:00",
                              "scan_mode": "local", **extra})
    db.update_proposal(pid, job_id=job_id)


def test_client_form_fields_are_cleaned_too(tmp_path):
    _task_for(tmp_path, "j-hostile", notes="pre\x1b[1mrelease\x00 <b>x</b>", path="/a\x0bb")
    text = _text(generator.generate({**JOB, "submitter": "acme"}, DIRTY[:1], "Formal Handover"))
    assert "prerelease <b>x</b>" in text and "/ab" in text
    assert "\x00" not in text and "\x1b" not in text


def test_long_notes_are_truncated(tmp_path):
    _task_for(tmp_path, "j-hostile", notes="N" * 50000)
    text = _text(generator.generate({**JOB, "submitter": "acme"}, DIRTY[:1], "Full Technical"))
    assert "[truncated]" in text and "N" * 2500 not in text


def test_report_comes_from_the_task_row_and_claims_no_approval(tmp_path):
    _task_for(tmp_path, "j-task", path="/shop", port=8443, notes="staging copy, avoid 02:00 batch")
    for template in ("Full Technical", "Formal Handover", "Executive Summary"):
        text = _text(generator.generate({**JOB, "submitter": "acme", "id": "j-task"}, DIRTY[:1], template))
        if template == "Executive Summary":
            continue   # no scope section in the business summary
        assert "2026-10-09 08:00 to 2026-10-09 16:30 UTC" in text
        assert "/shop" in text and "8443" in text and "staging copy, avoid 02:00 batch" in text
        assert "Target, path and port as requested by the client" in text
        assert "attested" not in text and "lead pentester" not in text
        for gone in ("Engagement purpose", "Division", "Rules of engagement", "Environment", "Out of scope"):
            assert gone not in text


def test_empty_task_fields_print_no_dash_rows(tmp_path):
    _task_for(tmp_path, "j-bare")
    text = _text(generator.generate({**JOB, "submitter": "acme", "id": "j-bare"}, DIRTY[:1], "Formal Handover"))
    assert "Client notes" not in text and "Path" not in text and "Port" not in text
