"""Detailed VAPT report builders (Varuna's own report design; no client template dependency).

Shared building blocks (cover, document control, scope/methodology, findings register,
per-finding pages, remediation roadmap, sign-off) composed into the client templates:
Full Technical, Formal Handover, Raw Findings, plus the section helpers the Executive Summary
reuses. python-docx only. Static contents list (not a Word TOC field) so the LibreOffice
docx->PDF delivery renders it; a TOC field would come out blank there.
"""
from __future__ import annotations

import datetime
import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "common"))
import models  # noqa: E402

NAVY = "0B2A43"
SEV_FILL = {"critical": "C0392B", "high": "E67E22", "medium": "F4D03F",
            "low": "3498DB", "info": "95A5A6"}
SEV_TEXT_DARK = {"medium"}   # yellow needs dark text
TIMEFRAME = {"critical": "Immediate (0-7 days)", "high": "Short term (within 30 days)",
             "medium": "Medium term (within 90 days)", "low": "Planned (next release cycle)",
             "info": "Planned (next release cycle)"}
TOOLS = [
    ("Katana", "Crawling and endpoint/form discovery of the in-scope application."),
    ("Nuclei", "Template-based detection of known vulnerabilities, exposures and misconfigurations."),
    ("SQLMap", "Automated detection of SQL injection on discovered parameters (safe profile: detection only)."),
]


# ---------------------------------------------------------------- low-level helpers
def _shade(cell, hex_fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _cell(cell, text, bold=False, fill=None, color=None, size=9.5, align=None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if align:
        p.alignment = align
    if fill:
        _shade(cell, fill)


def _table(doc, header, rows, widths=None, sev_col=None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        _cell(t.rows[0].cells[i], h, bold=True, fill=NAVY, color="FFFFFF")
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            if sev_col is not None and i == sev_col:
                s = str(v).lower()
                _cell(cells[i], str(v).capitalize(), bold=True, fill=SEV_FILL.get(s),
                      color="000000" if s in SEV_TEXT_DARK else "FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
            else:
                _cell(cells[i], v)
    if widths:
        from docx.shared import Cm
        t.autofit = False   # else LibreOffice (PDF delivery) ignores the widths
        for col, w in zip(t.columns, widths):   # gridCol widths: what LibreOffice actually reads
            col.width = Cm(w)
        for row in t.rows:
            for c, w in zip(row.cells, widths):
                c.width = Cm(w)
    doc.add_paragraph()
    return t


def _kv_table(doc, pairs):
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    for k, v in pairs:
        cells = t.add_row().cells
        _cell(cells[0], k, bold=True, fill="E8EEF3")
        _cell(cells[1], v if v not in (None, "") else "-")
    from docx.shared import Cm
    t.autofit = False
    t.columns[0].width, t.columns[1].width = Cm(4.2), Cm(12.3)
    for row in t.rows:
        row.cells[0].width, row.cells[1].width = Cm(4.2), Cm(12.3)
    doc.add_paragraph()


def _page_break(doc) -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _field(run, instr: str, cached: str = "1") -> None:
    """A complex field with a cached result: viewers that do not recalculate still show `cached`;
    LibreOffice (PDF delivery) and Word (on update) replace it with the real value."""
    def fld(kind):
        el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind); return el
    instr_el = OxmlElement("w:instrText"); instr_el.set(qn("xml:space"), "preserve"); instr_el.text = f" {instr} "
    text_el = OxmlElement("w:t"); text_el.text = cached
    for el in (fld("begin"), instr_el, fld("separate"), text_el, fld("end")):
        run._r.append(el)


def _footer(doc, label: str) -> None:
    p = doc.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"CONFIDENTIAL  |  {label}  |  Page ")
    r.font.size = Pt(8)
    r2 = p.add_run(); r2.font.size = Pt(8)
    _field(r2, "PAGE")


def new_document(footer_label: str) -> Document:
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(10.5)
    for lvl, size in ((1, 16), (2, 13), (3, 11.5)):
        h = doc.styles[f"Heading {lvl}"]
        h.font.name, h.font.size = "Calibri", Pt(size)
        h.font.color.rgb = RGBColor.from_string(NAVY)
    _footer(doc, footer_label)
    return doc


def _slug(text: str) -> str:
    return "sec_" + re.sub(r"[^a-z0-9]+", "_", text.lower()).strip("_")


def h1(doc, text: str):
    """Level-1 heading wrapped in a bookmark so the contents list can cite its page number."""
    para = doc.add_heading(text, level=1)
    name = _slug(text)
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(abs(hash(name)) % 100000))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), start.get(qn("w:id")))
    para._p.insert(1 if para._p.pPr is not None else 0, start)
    para._p.append(end)
    return para


# ---------------------------------------------------------------- data helpers
def counts_of(findings) -> dict:
    c = {s: 0 for s in models.SEVERITY_ORDER}
    for f in findings:
        c[(f.get("severity") or "info").lower()] = c.get((f.get("severity") or "info").lower(), 0) + 1
    return c


def overall_risk(counts: dict) -> str:
    for s in ("critical", "high", "medium", "low"):
        if counts.get(s):
            return s.capitalize()
    return "Informational" if counts.get("info") else "None identified"


def ordered(findings):
    """Most severe first, stable. Each gets a report ID (VAR-001...)."""
    out = sorted(findings, key=lambda f: models.severity_rank(f.get("severity", "")))
    return [dict(f, _id=f"VAR-{i:03d}") for i, f in enumerate(out, 1)]


def _flat(text) -> str:
    """Collapse tool-emitted hard line breaks so paragraphs wrap naturally."""
    return " ".join(str(text).split())


def _short(text, limit=170) -> str:
    """First sentence (capped) for table cells; the full text lives in the finding page."""
    t = _flat(text)
    cut = t.find(". ")
    t = t[:cut + 1] if 0 < cut < limit else t
    return t if len(t) <= limit else t[:limit].rstrip() + "..."



def _nm(f) -> str:
    """A finding can arrive without a name (manual entry left blank, or the Executive Summary's
    allow-list dropped an empty one); the layout must never assume it exists."""
    return f.get("name") or "Untitled finding"


def _refs(f) -> list[str]:
    refs = []
    cwe = (f.get("cwe") or "").upper().replace("CWE-", "").strip()
    if cwe.isdigit():
        refs.append(f"https://cwe.mitre.org/data/definitions/{cwe}.html")
    if f.get("cve"):
        refs.append(f"https://nvd.nist.gov/vuln/detail/{f['cve']}")
    if f.get("owasp"):
        refs.append("https://owasp.org/Top10/")
    return refs


# ---------------------------------------------------------------- sections
def cover(doc, title: str, subtitle: str, job: dict, ctx: dict) -> None:
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("VARUNA"); r.bold = True; r.font.size = Pt(14); r.font.color.rgb = RGBColor.from_string("0B5FA5")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title); r.bold = True; r.font.size = Pt(28); r.font.color.rgb = RGBColor.from_string(NAVY)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(subtitle); r.font.size = Pt(13)
    doc.add_paragraph()
    _kv_table(doc, [
        ("Client", ctx.get("client")), ("Target", job.get("target", "")),
        ("Engagement purpose", ctx.get("purpose")), ("Division", ctx.get("division")),
        ("Report date", datetime.date.today().isoformat()),
        ("Reference ID", job.get("id", "")), ("Classification", "CONFIDENTIAL"),
    ])
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("This document contains sensitive security information intended only for the named recipient. "
                  "Do not distribute without written authorization.")
    r.italic = True; r.font.size = Pt(9)
    _page_break(doc)


def document_control(doc, status: str) -> None:
    h1(doc, "Document Control")
    _table(doc, ["Version", "Date", "Status", "Description"],
           [["1.0", datetime.date.today().isoformat(), status, "Initial report generated by Varuna"]],
           widths=[2, 3, 4, 7.5])


def contents(doc, items: list[str]) -> None:
    from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
    from docx.shared import Cm
    doc.add_heading("Contents", level=1)
    for i, name in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(Cm(16.3), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        p.add_run(f"{i}.  {name}\t")
        _field(p.add_run(), f"PAGEREF {_slug(name)} \\h", cached="-")
    _page_break(doc)


def executive_overview(doc, findings, counts, job, ctx, with_actions=True) -> None:
    """Business-safe: uses only name/severity/risk/impact/remediation (no tool or payload detail)."""
    h1(doc, "Executive Summary")
    risk = overall_risk(counts)
    doc.add_paragraph(
        f"Varuna performed an automated web application vulnerability assessment of "
        f"{job.get('target', 'the target')} for {ctx.get('client') or 'the client'}. "
        f"The assessment identified {len(findings)} confirmed finding(s). "
        f"The overall risk rating is {risk.upper()}."
    )
    t = doc.add_table(rows=2, cols=len(models.SEVERITY_ORDER))
    t.style = "Table Grid"
    for i, s in enumerate(models.SEVERITY_ORDER):
        _cell(t.rows[0].cells[i], s.capitalize(), bold=True, fill=SEV_FILL[s],
              color="000000" if s in SEV_TEXT_DARK else "FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
        _cell(t.rows[1].cells[i], counts.get(s, 0), bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph()
    if counts["critical"] or counts["high"]:
        doc.add_paragraph(
            "Critical and high severity findings can lead to data exposure or system compromise and "
            "should be remediated first. See the Remediation Roadmap for recommended timeframes."
        )
    doc.add_heading("Key Findings", level=2)
    for f in findings[:5]:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(f"{_nm(f)} ({f.get('severity', '').capitalize()}). ").bold = True
        if f.get("impact"):
            p.add_run(f["impact"] + " ")
        if with_actions and f.get("remediation"):
            p.add_run("Action: " + f["remediation"])


def scope_methodology(doc, job, ctx, with_tools=True) -> None:
    h1(doc, "Scope and Methodology")
    doc.add_heading("Engagement Details", level=2)
    roe = ctx.get("roe") or {}
    roe_txt = ", ".join(k.replace("_", " ") for k, v in roe.items() if v) or "Standard safe profile"
    _kv_table(doc, [
        ("Target", job.get("target", "")), ("In scope", ctx.get("in_scope") or job.get("target", "")),
        ("Out of scope", ctx.get("out_of_scope") or "None declared"),
        ("Environment", ctx.get("environment")), ("Test window", ctx.get("test_window")),
        ("Rules of engagement", roe_txt), ("Testing approach", "Automated, non-destructive (safe profile)"),
        ("Authorization", "Client attested ownership / authorization to test; approved by the lead pentester."),
    ])
    doc.add_heading("Approach", level=2)
    for step in (
        "Discovery: the application was crawled to enumerate pages, parameters and forms.",
        "Detection: known-vulnerability and misconfiguration checks, and injection testing, were run against what was found.",
        "Correlation: duplicate and related results were merged and each finding was classified against OWASP Top 10 (2021) and CWE.",
        "Triage: results were reviewed by the security team; false positives were removed before this report.",
        "Reporting: each finding is rated, described, evidenced, and paired with remediation guidance.",
    ):
        doc.add_paragraph(step, style="List Number")
    if with_tools:
        doc.add_heading("Tooling", level=2)
        _table(doc, ["Tool", "Purpose"], [list(t) for t in TOOLS], widths=[3.5, 13])


def coverage(doc, title: str, paras: list[str]) -> None:
    h1(doc, title)
    for para in paras:
        doc.add_paragraph(para)


def rating_methodology(doc) -> None:
    h1(doc, "Risk Rating Methodology")
    doc.add_paragraph("Findings are rated by severity, informed by CVSS where a score is available, "
                      "and by exploitability and business impact.")
    _table(doc, ["Severity", "Meaning", "Target response"], [
        ["critical", "Trivially exploitable; leads to full compromise or bulk data loss.", TIMEFRAME["critical"]],
        ["high", "Exploitable with moderate effort; significant data or access exposure.", TIMEFRAME["high"]],
        ["medium", "Requires specific conditions or chaining; limited impact.", TIMEFRAME["medium"]],
        ["low", "Minor weakness; hardening opportunity.", TIMEFRAME["low"]],
        ["info", "No direct risk; observation or best-practice note.", TIMEFRAME["info"]],
    ], widths=[2.5, 9.5, 4.5], sev_col=0)


def findings_register(doc, findings, show_verdict=False) -> None:
    h1(doc, "Findings Summary")
    if not findings:
        doc.add_paragraph("No findings were identified within the automated scope of this assessment.")
        return
    header = ["ID", "Finding", "Severity", "Affected asset", "CVSS", "OWASP"]
    if show_verdict:
        header.append("Verdict")
    rows = []
    for f in findings:
        row = [f["_id"], _nm(f), f.get("severity", "info"), f.get("host") or f.get("url") or "-",
               f.get("cvss") if f.get("cvss") is not None else "-", f.get("owasp") or "-"]
        if show_verdict:
            row.append((f.get("verdict") or "tp").upper())
        rows.append(row)
    _table(doc, header, rows, sev_col=2,
           widths=[2.0, 4.2, 2.4, 3.1, 1.3, 3.5] + ([1.5] if show_verdict else []))


def finding_detail(doc, f, evidence=True, show_verdict=False) -> None:
    sev = (f.get("severity") or "info").lower()
    doc.add_heading(f"{f['_id']}  {_nm(f)}", level=2)
    t = doc.add_table(rows=1, cols=4)
    t.style = "Table Grid"
    _cell(t.rows[0].cells[0], "Severity", bold=True, fill="E8EEF3")
    _cell(t.rows[0].cells[1], sev.capitalize(), bold=True, fill=SEV_FILL.get(sev),
          color="000000" if sev in SEV_TEXT_DARK else "FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
    _cell(t.rows[0].cells[2], "Status", bold=True, fill="E8EEF3")
    _cell(t.rows[0].cells[3], (f.get("status") or "open").capitalize())
    doc.add_paragraph()
    pairs = [("Affected asset", f.get("host")), ("URL / parameter", f.get("url")),
             ("CVSS", f.get("cvss")), ("CWE", f.get("cwe")), ("CVE", f.get("cve")),
             ("OWASP Top 10", f.get("owasp")), ("Risk rating", f.get("risk_rating"))]
    if evidence:
        pairs.append(("Detected by", (f.get("tool") or "").capitalize()))
    if show_verdict:
        pairs.append(("Triage verdict", "True positive" if (f.get("verdict") or "tp") == "tp" else "False positive"))
    _kv_table(doc, [(k, v) for k, v in pairs if v not in (None, "")])
    for title, key in (("Description", "description"), ("Impact", "impact"), ("Recommendation", "remediation")):
        if f.get(key):
            doc.add_heading(title, level=3)
            doc.add_paragraph(_flat(f[key]))
    if evidence and f.get("evidence"):
        doc.add_heading("Evidence", level=3)
        p = doc.add_paragraph()
        r = p.add_run(str(f["evidence"])); r.font.name = "Courier New"; r.font.size = Pt(8.5)
    refs = _refs(f)
    if refs:
        doc.add_heading("References", level=3)
        for ref in refs:
            doc.add_paragraph(ref, style="List Bullet")
    doc.add_heading("Retest", level=3)
    doc.add_paragraph("After the fix is deployed, request a retest to confirm closure before marking this finding resolved.")


def roadmap(doc, findings) -> None:
    h1(doc, "Remediation Roadmap")
    if not findings:
        doc.add_paragraph("No remediation actions are required from this assessment.")
        return
    doc.add_paragraph("Recommended order of work, grouped by target timeframe.")
    rows = [[TIMEFRAME[(f.get("severity") or "info").lower()] if (f.get("severity") or "info").lower() in TIMEFRAME else TIMEFRAME["info"],
             f["_id"], _nm(f), (f.get("severity") or "info"),
             _short(f["remediation"]) if f.get("remediation") else "See finding detail."]
            for f in findings]
    _table(doc, ["Timeframe", "ID", "Finding", "Severity", "Action"], rows, widths=[3.0, 2.2, 3.8, 2.0, 5.5], sev_col=3)


def signoff(doc, roles=("Prepared by (Pentester)", "Reviewed by (Lead Pentester)",
                        "Approved by (Governance)", "Received by (Client)")) -> None:
    h1(doc, "Sign-off and Acceptance")
    doc.add_paragraph("By signing, the parties acknowledge issue and receipt of this report.")
    _table(doc, ["Role", "Name", "Signature", "Date"], [[r, "", "", ""] for r in roles], widths=[5.5, 4, 4, 3])
    doc.add_paragraph()


def appendix_glossary(doc) -> None:
    h1(doc, "Appendix: Glossary")
    _table(doc, ["Term", "Meaning"], [
        ["VAPT", "Vulnerability Assessment and Penetration Testing."],
        ["CVSS", "Common Vulnerability Scoring System; a 0-10 severity score."],
        ["CWE", "Common Weakness Enumeration; a catalogue of software weakness types."],
        ["CVE", "Common Vulnerabilities and Exposures; identifier for a publicly known flaw."],
        ["OWASP Top 10", "Awareness list of the most critical web application security risks."],
        ["False positive", "A reported issue that triage found not to be a real vulnerability."],
    ], widths=[3.5, 13])


# ---------------------------------------------------------------- templates
def _clean(findings):
    return ordered([f for f in findings if (f.get("verdict") or "tp") != "fp"])


def full_technical(job, findings, ctx, coverage_paras, coverage_title) -> Document:
    fs = _clean(findings)
    doc = new_document("Full Technical Report")
    cover(doc, "Web Application VAPT Report", "Full Technical Report", job, ctx)
    document_control(doc, "Draft for review")
    contents(doc, ["Executive Summary", "Scope and Methodology", coverage_title, "Risk Rating Methodology",
                   "Findings Summary", "Detailed Findings", "Remediation Roadmap", "Appendix: Glossary"])
    c = counts_of(fs)
    executive_overview(doc, fs, c, job, ctx)
    scope_methodology(doc, job, ctx)
    coverage(doc, coverage_title, coverage_paras)
    rating_methodology(doc)
    findings_register(doc, fs)
    h1(doc, "Detailed Findings")
    for f in fs:
        finding_detail(doc, f)
    roadmap(doc, fs)
    appendix_glossary(doc)
    return doc


def formal_handover(job, findings, ctx, coverage_paras, coverage_title) -> Document:
    fs = _clean(findings)
    doc = new_document("Formal Handover Report")
    cover(doc, "Web Application VAPT Report", "Formal Handover", job, ctx)
    document_control(doc, "Final")
    contents(doc, ["Executive Summary", "Scope and Methodology", coverage_title, "Findings Summary",
                   "Remediation Roadmap", "Sign-off and Acceptance"])
    c = counts_of(fs)
    executive_overview(doc, fs, c, job, ctx)
    scope_methodology(doc, job, ctx, with_tools=False)
    coverage(doc, coverage_title, coverage_paras)
    findings_register(doc, fs)
    roadmap(doc, fs)
    signoff(doc)
    return doc


def raw_findings(job, findings, ctx, coverage_paras, coverage_title) -> Document:
    fs = ordered(findings)   # everything, including false positives, marked
    doc = new_document("Raw Findings")
    cover(doc, "Web Application VAPT Report", "Raw Findings (lightly processed, with triage marks)", job, ctx)
    c = counts_of(fs)
    tp = sum(1 for f in fs if (f.get("verdict") or "tp") == "tp")
    h1(doc, "Overview")
    doc.add_paragraph(f"{len(fs)} result(s): {tp} confirmed (TP), {len(fs) - tp} marked false positive (FP). "
                      f"Severity before triage: " + ", ".join(f"{s} {c[s]}" for s in models.SEVERITY_ORDER) + ".")
    findings_register(doc, fs, show_verdict=True)
    h1(doc, "Results")
    for f in fs:
        finding_detail(doc, f, show_verdict=True)
    coverage(doc, coverage_title, coverage_paras)
    return doc
