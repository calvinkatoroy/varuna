"""Content + live findings -> resolved sections of plain items. The same items become the docx (python-docx runs
take text literally: nothing here is ever interpreted as markup) and the preview JSON of the audit page."""
from __future__ import annotations

import io
import os
import sys

from docx.enum.text import WD_ALIGN_PARAGRAPH

sys.path.insert(0, os.path.dirname(__file__))
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "common"))
import detailed  # noqa: E402
import generator  # noqa: E402
import models  # noqa: E402

_PREVIEW_KEYS = ("id", "_id", "name", "severity", "host", "url", "impact", "remediation", "description",
                 "cve", "cwe", "tool", "verdict", "evidence")
_NONE = "No findings were identified within the automated scope of this assessment."


def _ctx(task: dict, org_name: str) -> dict:
    window = f"{generator._when(task.get('not_before'))} to {generator._when(task.get('not_after'))} UTC" \
        if task.get("not_after") else ""
    return generator._clean_value({
        "client": org_name, "test_window": window, "path": task.get("path"),
        "port": str(task["port"]) if task.get("port") else "", "notes": task.get("notes"),
        "scan_mode": {"local": "Local (agent on the client's network)", "cloud": "Cloud (public target)"}.get(
            task.get("scan_mode"), "")})


def _sev(f: dict) -> str:
    s = (f.get("severity") or "info").lower()
    return s if s in detailed.TIMEFRAME else "info"


def resolve(content: dict, task: dict, org_name: str, findings: list[dict]) -> list[dict]:
    live = detailed.ordered([f for f in generator._clean_value(findings) if (f.get("verdict") or "tp") != "fp"])
    overrides: dict = {}
    for sec in content["sections"]:
        for b in sec["blocks"]:
            if b["type"] == "finding_overrides":
                overrides = b.get("overrides") or {}
    merged = []
    for f in live:
        o = overrides.get(f["id"]) or {}
        merged.append({**f, **{k: generator._clean_text(v, 4000) for k, v in o.items()
                               if k in ("impact", "remediation") and isinstance(v, str)}})
    counts = detailed.counts_of(merged)
    ctx = _ctx(task, org_name)
    out = []
    for sec in content["sections"]:
        items: list[dict] = []
        for b in sec["blocks"]:
            t = b["type"]
            if t in ("paragraph", "bullet"):
                items.append({"kind": t, "text": generator._clean_text(b.get("text", ""), 4000)})
            elif t == "cover":
                items.append({"kind": "cover", "client": org_name, "target": task["target"], "ref": task["id"]})
            elif t == "severity_table":
                items.append({"kind": "counts", "counts": counts, "risk": detailed.overall_risk(counts)})
            elif t == "scope_table":
                rows = [("Target", task["target"]), ("Path", ctx.get("path")), ("Port", ctx.get("port")),
                        ("Scan mode", ctx.get("scan_mode")), ("Test window", ctx.get("test_window")),
                        ("Client notes", ctx.get("notes")),
                        ("Testing approach", "Automated, non-destructive (safe profile)"),
                        ("Scope", "Target, path and port as requested by the client")]
                items.append({"kind": "kv", "rows": [[k, v] for k, v in rows if v]})
            elif t == "tools_table":
                items.append({"kind": "table", "header": ["Tool", "Purpose"], "rows": [list(x) for x in detailed.TOOLS],
                              "sev_col": None, "widths": [3.5, 13]})
            elif t == "findings_table":
                if not merged:
                    items.append({"kind": "paragraph", "text": _NONE})
                else:
                    items.append({"kind": "table", "header": ["ID", "Finding", "Severity", "Affected asset", "CVSS", "OWASP"],
                                  "rows": [[f["_id"], f.get("name") or "Untitled finding", f.get("severity") or "info",
                                            f.get("host") or f.get("url") or "-",
                                            f["cvss"] if f.get("cvss") is not None else "-", f.get("owasp") or "-"]
                                           for f in merged], "sev_col": 2, "widths": [2.0, 4.2, 2.4, 3.1, 1.3, 3.5]})
            elif t == "roadmap":
                if not merged:
                    items.append({"kind": "paragraph", "text": "No remediation actions are required from this assessment."})
                else:
                    items.append({"kind": "table", "header": ["Timeframe", "ID", "Finding", "Severity", "Action"],
                                  "rows": [[detailed.TIMEFRAME[_sev(f)], f["_id"], f.get("name") or "Untitled finding",
                                            f.get("severity") or "info",
                                            detailed._short(f["remediation"]) if f.get("remediation") else "See finding detail."]
                                           for f in merged], "sev_col": 3, "widths": [3.0, 2.2, 3.8, 2.0, 5.5]})
            elif t == "finding_overrides":
                items.extend({"kind": "finding", "f": f} for f in merged)
        out.append({"id": sec["id"], "title": sec["title"], "items": items})
    return out


def _counts_table(doc, it: dict) -> None:
    c = it["counts"]
    t = doc.add_table(rows=2, cols=len(models.SEVERITY_ORDER))
    t.style = "Table Grid"
    for i, s in enumerate(models.SEVERITY_ORDER):
        detailed._cell(t.rows[0].cells[i], s.capitalize(), bold=True, fill=detailed.SEV_FILL[s],
                       color="000000" if s in detailed.SEV_TEXT_DARK else "FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
        detailed._cell(t.rows[1].cells[i], c.get(s, 0), bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph()
    doc.add_paragraph(f"Overall risk rating: {it['risk'].upper()}")


def _emit(doc, it: dict) -> None:
    k = it["kind"]
    if k == "paragraph":
        doc.add_paragraph(it["text"])
    elif k == "bullet":
        doc.add_paragraph(it["text"], style="List Bullet")
    elif k == "kv":
        detailed._kv_table(doc, [tuple(r) for r in it["rows"]])
    elif k == "table":
        detailed._table(doc, it["header"], it["rows"], widths=it.get("widths"), sev_col=it.get("sev_col"))
    elif k == "counts":
        _counts_table(doc, it)
    elif k == "finding":
        detailed.finding_detail(doc, it["f"])


def render_docx(content: dict, task: dict, org_name: str, findings: list[dict]) -> bytes:
    sections = resolve(content, task, org_name, findings)
    doc = detailed.new_document("Penetration Test Report")
    ctx = _ctx(task, org_name)
    titles = [s["title"] for s in sections if s["id"] != "cover"]
    for s in sections:
        if s["id"] == "cover":
            for it in s["items"]:
                detailed.cover(doc, "Web Application VAPT Report", "Penetration Test Report",
                               {"target": it["target"], "id": it["ref"]}, ctx)
            detailed.contents(doc, titles)
            continue
        detailed.h1(doc, s["title"])
        for it in s["items"]:
            _emit(doc, it)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def preview(content: dict, task: dict, org_name: str, findings: list[dict]) -> list[dict]:
    out = resolve(content, task, org_name, findings)
    for s in out:
        for it in s["items"]:
            if it["kind"] == "finding":
                f = {k: it["f"].get(k) for k in _PREVIEW_KEYS}
                f["evidence"] = (f["evidence"] or "")[:500]
                it["f"] = f
    return out
