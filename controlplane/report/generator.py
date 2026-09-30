"""`.docx` report generation via python-docx (SRS §4.8, REQ-44 to REQ-50).

Executive Summary, Full Technical, Formal Handover and Raw Findings use the detailed layout in
detailed.py (cover, document control, scope/methodology, findings register, per-finding pages,
remediation roadmap, sign-off). OWASP Web App and ILCS Internal are the older simple layouts. The Executive Summary is built ONLY from
sanitized findings (sanitize.py, REQ-49); Full Technical carries full detail.
"""
from __future__ import annotations

import datetime
import io
import os
import sys

from docx import Document

sys.path.insert(0, os.path.dirname(__file__))                                    # sanitize
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "common"))      # models

import db  # noqa: E402
import detailed  # noqa: E402
import models  # noqa: E402
import sanitize  # noqa: E402

TEMPLATES = ("Full Technical", "Formal Handover", "Executive Summary", "Raw Findings",
             "OWASP Web App", "ILCS Internal")


def _ctx(job: dict) -> dict:
    """Engagement details from the originating proposal (scope, environment, RoE...)."""
    try:
        p = db.get_proposal_by_job(job.get("id") or "") or {}
    except Exception:   # no DB (offline unit use): a report must still generate
        p = {}
    return {"client": job.get("submitter") or p.get("submitter"), "purpose": p.get("purpose"),
            "division": p.get("division"), "environment": p.get("environment"),
            "test_window": p.get("test_window"), "in_scope": p.get("in_scope"),
            "out_of_scope": p.get("out_of_scope"), "roe": p.get("roe")}

# Coverage disclaimer: every report MUST carry it so no reader mistakes "no findings" for
# "fully secure" (SRS §1.4 Coverage and Limitations). Automated scanning covers only part
# of the OWASP Top 10; the rest needs a human pentester.
COVERAGE_TITLE = "Scope and Coverage"
COVERAGE_PARAS = [
    "This report reflects an AUTOMATED web application VAPT scan. Automated scanning covers "
    "only part of the OWASP Top 10 2021: strong coverage of A03 (Injection), A05 (Security "
    "Misconfiguration), and A06 (Vulnerable and Outdated Components); partial coverage of A02 "
    "(Cryptographic Failures) and A10 (Server-Side Request Forgery).",
    "The following categories are NOT assessed by this automated scan and require manual "
    "penetration testing: A01 (Broken Access Control), A04 (Insecure Design and business "
    "logic), A07 (Identification and Authentication Failures), and A09 (Security Logging and "
    "Monitoring Failures).",
    "The absence of findings does NOT mean the target is fully secure; it means no issues were "
    "detected within the automated scope above. A manual penetration test is recommended to "
    "cover the categories automated tools cannot reach, followed by re-testing after remediation.",
]


def _add_coverage(doc: Document) -> None:
    doc.add_heading(COVERAGE_TITLE, level=2)
    for para in COVERAGE_PARAS:
        doc.add_paragraph(para)


def _severity_counts(findings: list[dict]) -> dict:
    counts = {s: 0 for s in models.SEVERITY_ORDER}
    for f in findings:
        s = (f.get("severity") or "info").lower()
        counts[s] = counts.get(s, 0) + 1
    return counts


def _add_meta(doc: Document, job: dict) -> None:
    t = doc.add_table(rows=0, cols=2)
    for label, value in (
        ("Target", job.get("target", "")),
        ("Date", datetime.date.today().isoformat()),
        ("Reference ID", job.get("id", "")),
    ):
        row = t.add_row().cells
        row[0].text, row[1].text = label, str(value)


def _add_severity_table(doc: Document, findings: list[dict]) -> None:
    counts = _severity_counts(findings)
    doc.add_heading("Severity Summary", level=2)
    t = doc.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Severity", "Count"
    for sev in models.SEVERITY_ORDER:
        row = t.add_row().cells
        row[0].text, row[1].text = sev.capitalize(), str(counts.get(sev, 0))


def _executive_summary(job: dict, findings: list[dict]) -> Document:
    ctx = _ctx(job)
    # Allow-list first (REQ-49): technical fields never reach this document. Only then add IDs.
    fs = detailed.ordered(sanitize.sanitize_findings(
        [f for f in findings if (f.get("verdict") or "tp") != "fp"]))
    doc = detailed.new_document("Executive Summary")
    detailed.cover(doc, "Web Application VAPT Report", "Executive Summary", job, ctx)
    detailed.executive_overview(doc, fs, detailed.counts_of(fs), job, ctx)
    detailed.roadmap(doc, fs)
    detailed.coverage(doc, COVERAGE_TITLE, COVERAGE_PARAS)
    doc.add_heading("Recommended Next Steps", level=1)
    doc.add_paragraph(
        "Prioritize remediation of critical and high findings, re-test after fixes, and "
        "schedule a follow-up assessment to confirm closure."
    )
    return doc


def _add_evidence(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Courier New"   # monospace evidence block (mirrors REQ-42)


def _add_finding_detail(doc: Document, f: dict, level: int = 3) -> None:
    doc.add_heading(f.get("name", "Finding"), level=level)
    for label, key in (
        ("Severity", "severity"), ("Host", "host"), ("URL", "url"),
        ("CVE", "cve"), ("CVSS", "cvss"), ("CWE", "cwe"), ("OWASP", "owasp"),
        ("Description", "description"), ("Impact", "impact"), ("Remediation", "remediation"),
    ):
        if f.get(key):
            doc.add_paragraph(f"{label}: {f[key]}")
    if f.get("evidence"):
        doc.add_paragraph("Evidence:")
        _add_evidence(doc, f["evidence"])


def _full_technical(job: dict, findings: list[dict]) -> Document:
    return detailed.full_technical(job, findings, _ctx(job), COVERAGE_PARAS, COVERAGE_TITLE)


def _formal_handover(job: dict, findings: list[dict]) -> Document:
    return detailed.formal_handover(job, findings, _ctx(job), COVERAGE_PARAS, COVERAGE_TITLE)


def _raw_findings(job: dict, findings: list[dict]) -> Document:
    return detailed.raw_findings(job, findings, _ctx(job), COVERAGE_PARAS, COVERAGE_TITLE)


def _owasp_web_app(job: dict, findings: list[dict]) -> Document:
    doc = Document()
    doc.add_heading("OWASP Web Application Report, Web VAPT", level=0)
    _add_meta(doc, job)
    _add_coverage(doc)
    _add_severity_table(doc, findings)
    doc.add_heading("Findings by OWASP Category", level=2)
    groups: dict = {}
    for f in findings:
        groups.setdefault(f.get("owasp") or "Uncategorized", []).append(f)
    for category in sorted(groups):
        doc.add_heading(category, level=3)
        for f in groups[category]:
            _add_finding_detail(doc, f, level=4)
    return doc


# --- ILCS Internal template: Bahasa Indonesia + sign-off table (BR-2, REQ-81, TBD-1) ---

COVERAGE_PARAS_ID = [
    "Laporan ini merupakan hasil pemindaian VAPT aplikasi web secara OTOMATIS. Pemindaian "
    "otomatis hanya mencakup sebagian OWASP Top 10 2021: cakupan kuat pada A03 (Injection), "
    "A05 (Kesalahan Konfigurasi Keamanan), dan A06 (Komponen Rentan dan Usang); cakupan "
    "sebagian pada A02 (Kegagalan Kriptografi) dan A10 (Server-Side Request Forgery).",
    "Kategori berikut TIDAK diuji oleh pemindaian otomatis ini dan memerlukan pengujian "
    "penetrasi manual: A01 (Broken Access Control), A04 (Desain Tidak Aman dan logika bisnis), "
    "A07 (Kegagalan Identifikasi dan Autentikasi), dan A09 (Kegagalan Pencatatan dan Pemantauan).",
    "Tidak adanya temuan TIDAK berarti target sepenuhnya aman; artinya tidak ada masalah yang "
    "terdeteksi dalam cakupan otomatis di atas. Pengujian penetrasi manual disarankan untuk "
    "menutup kategori yang tidak terjangkau alat otomatis, diikuti pengujian ulang setelah remediasi.",
]

_SEV_ID = {"critical": "Kritis", "high": "Tinggi", "medium": "Sedang", "low": "Rendah", "info": "Informasi"}


def _ilcs_internal(job: dict, findings: list[dict]) -> Document:
    doc = Document()
    doc.add_heading("Laporan VAPT Aplikasi Web, ILCS", level=0)

    meta = doc.add_table(rows=0, cols=2)
    for label, value in (("Target", job.get("target", "")),
                         ("Tanggal", datetime.date.today().isoformat()),
                         ("ID Referensi", job.get("id", ""))):
        row = meta.add_row().cells
        row[0].text, row[1].text = label, str(value)

    doc.add_heading("Ruang Lingkup dan Cakupan", level=2)
    for para in COVERAGE_PARAS_ID:
        doc.add_paragraph(para)

    counts = _severity_counts(findings)
    doc.add_heading("Ringkasan Tingkat Keparahan", level=2)
    tbl = doc.add_table(rows=1, cols=2)
    tbl.rows[0].cells[0].text, tbl.rows[0].cells[1].text = "Tingkat", "Jumlah"
    for sev in models.SEVERITY_ORDER:
        row = tbl.add_row().cells
        row[0].text, row[1].text = _SEV_ID[sev], str(counts.get(sev, 0))

    doc.add_heading("Temuan", level=2)
    for f in sorted(findings, key=lambda x: models.severity_rank(x.get("severity", ""))):
        doc.add_heading(f.get("name", "Temuan"), level=3)
        for label, key in (("Tingkat", "severity"), ("Host", "host"), ("URL", "url"),
                           ("CVE", "cve"), ("CVSS", "cvss"), ("CWE", "cwe"), ("OWASP", "owasp"),
                           ("Deskripsi", "description"), ("Dampak", "impact"), ("Rekomendasi", "remediation")):
            if f.get(key):
                doc.add_paragraph(f"{label}: {f[key]}")
        if f.get("evidence"):
            doc.add_paragraph("Bukti:")
            _add_evidence(doc, f["evidence"])

    doc.add_heading("Persetujuan", level=2)
    sign = doc.add_table(rows=1, cols=4)
    for i, head in enumerate(("Nama", "Jabatan", "Tanda Tangan", "Tanggal")):
        sign.rows[0].cells[i].text = head
    for _ in range(2):
        sign.add_row()   # blank rows for signatures
    return doc


_BUILDERS = {
    "Executive Summary": _executive_summary,
    "Full Technical": _full_technical,
    "Formal Handover": _formal_handover,
    "Raw Findings": _raw_findings,
    "OWASP Web App": _owasp_web_app,
    "ILCS Internal": _ilcs_internal,
}


def generate(job: dict, findings: list[dict], template: str) -> bytes:
    if template not in _BUILDERS:
        raise ValueError(f"unknown template: {template}")
    doc = _BUILDERS[template](job, findings)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
